#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""dp.py - công cụ CƠ HỌC cho skill dich-pro. Script KHÔNG dịch: chỉ tách / hiển thị / ghi / kiểm tra / ghép.
Lệnh: prep | peek | next | add | fix | check | sample | merge   (chi tiết trong SKILL.md)
Chạy trong thư mục chứa file cần dịch; dữ liệu làm việc nằm ở ./dp_work/"""
import sys, os, re, json, random, argparse, shutil
from collections import Counter

for _s in (sys.stdout, sys.stdin):
    try: _s.reconfigure(encoding='utf-8')
    except Exception: pass

W = 'dp_work'
P = lambda n: os.path.join(W, n)
CJK = re.compile(r'[\uac00-\ud7a3\u1100-\u11ff\u3130-\u318f\u3000-\u303f\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uff00-\uffef]')
CODE0 = r'\\[nrt]|%\d*[sdfx]|\{\d+\}|\$\d+|\^\d|#[0-9A-Za-z]|</?[A-Za-z][^<>]{0,20}>|[|@~]'
NLE, TABE = '¶', '⇥'
PR = re.compile(r'^(\s*[^\s#!=:][^=:]*?\s*[=:]\s*)(.*)$')
HDR = dict(
    src=('src', 'source', 'gốc', 'goc', 'text', 'original', 'orig', 'zh', 'cn', 'en', '文本', '原文', 'bản gốc'),
    dst=('dst', 'dịch', 'dich', 'vi', 'vn', 'translation', 'trans', 'target', '译文', 'bản dịch', 'ban dich'),
    ctx=('ctx', 'context', 'ghi chú', 'ghi chu', 'note', 'notes', 'speaker', 'người nói', 'nhân vật'),
    lim=('limit', 'max', 'maxlen', 'len', 'giới hạn', 'gioi han', 'độ dài', 'do dai'))
HDR_ALL = {n for v in HDR.values() for n in v} | {'id', 'key', 'stt', 'idx', 'index', 'mã'}


def die(m):
    print('LỖI:', m); sys.exit(1)


def jl(n, d=None):
    p = P(n)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else d


def jd(n, o):
    json.dump(o, open(P(n), 'w', encoding='utf-8'), ensure_ascii=False)


def rd(p, enc):
    b = open(p, 'rb').read()
    bom = enc.lower() in ('utf-8', 'utf8') and b.startswith(b'\xef\xbb\xbf')
    return b.decode('utf-8-sig' if enc.lower() in ('utf-8', 'utf8') else enc), bom


def wr(p, t, enc, bom):
    os.makedirs(os.path.dirname(p) or '.', exist_ok=True)
    try:
        b = t.encode(enc)
    except UnicodeEncodeError as e:
        die(f'encoding {enc} không chứa được ký tự: {e.object[e.start:e.end]!r}')
    open(p, 'wb').write((b'\xef\xbb\xbf' if bom else b'') + b)


def colmap(first, ncols, a):
    low = [str(c).strip().lower() for c in first]
    has = any(c in HDR_ALL for c in low)
    cols = {}
    if has:
        for f, names in HDR.items():
            for i, c in enumerate(low):
                if c in names: cols[f] = i; break
    for f in HDR:
        if getattr(a, f, None) is not None: cols[f] = getattr(a, f)
    if not has:
        if 'src' not in cols:
            cols['src'] = 0 if ncols == 1 else 1
        if 'dst' not in cols and ncols >= 3:
            cols['dst'] = 2
    if 'src' not in cols:
        die('không nhận ra cột gốc; dùng --src N (đếm từ 0)')
    if 'dst' not in cols:
        cols['dst'] = ncols; cols['dst_new'] = True
    return has, cols


def lim_of(x):
    m = re.search(r'\d+', str(x or ''))
    return int(m.group()) if m else None


def load_tsv(p, a, prop=False):
    t, bom = rd(p, a.enc)
    nl = '\r\n' if '\r\n' in t else '\n'
    lines = t.split(nl)
    tail = bool(lines) and lines[-1] == ''
    if tail: lines.pop()
    rows = []
    meta = dict(fmt='prop' if prop else 'tsv', nl=nl, tail=tail, bom=bom, hdr=False, cols={})
    if prop:
        for i, l in enumerate(lines):
            m = PR.match(l)
            if m and m.group(2).strip(): rows.append((i, m.group(2), '', None))
        return meta, rows, lines
    nc = max([len(l.split('\t')) for l in lines[:50]] or [1])
    has, cols = colmap(lines[0].split('\t') if lines else [], nc, a)
    meta.update(hdr=has, cols=cols)
    for i in range(1 if has else 0, len(lines)):
        c = lines[i].split('\t')
        if len(c) > cols['src'] and c[cols['src']].strip():
            ctx = c[cols['ctx']] if 'ctx' in cols and len(c) > cols['ctx'] else ''
            lim = lim_of(c[cols['lim']]) if 'lim' in cols and len(c) > cols['lim'] else None
            rows.append((i, c[cols['src']], ctx, lim))
    return meta, rows, lines


def load_xlsx(p, a):
    import openpyxl
    wb = openpyxl.load_workbook(p)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    has, cols = colmap([c.value or '' for c in ws[1]], ws.max_column, a)
    meta = dict(fmt='xlsx', hdr=has, cols=cols, bom=False)
    rows = []
    for r in range(2 if has else 1, ws.max_row + 1):
        v = ws.cell(r, cols['src'] + 1).value
        if v is None or not str(v).strip(): continue
        ctx = ws.cell(r, cols['ctx'] + 1).value if 'ctx' in cols else ''
        lim = lim_of(ws.cell(r, cols['lim'] + 1).value) if 'lim' in cols else None
        rows.append((r, str(v), str(ctx or ''), lim))
    return meta, rows, (wb, ws)


def load_json(p, a):
    t, bom = rd(p, a.enc)
    d = json.loads(t)
    if not isinstance(d, dict):
        die('JSON chỉ hỗ trợ dạng {id: text}; dạng khác hãy đổi sang TSV rồi prep lại')
    rows = [(k, v, '', None) for k, v in d.items() if isinstance(v, str) and v.strip()]
    return dict(fmt='json', bom=bom, indent=2 if '\n' in t else None, hdr=False, cols={}), rows, d


def loader(p):
    return {'.xlsx': load_xlsx, '.xlsm': load_xlsx, '.json': load_json,
            '.properties': lambda p, a: load_tsv(p, a, True)}.get(os.path.splitext(p)[1].lower(), load_tsv)


def skip_it(s, lang):
    t = s.strip()
    if not t or not re.search(r'[^\W\d_]', t): return True
    if CJK.search(t): return False
    if re.fullmatch(r'[\w./\\-]+', t):
        return lang == 'zh' or bool(re.search(r'[_./\\\d]', t))
    return False


def esc(s):
    return s.replace('\r\n', NLE).replace('\n', NLE).replace('\t', TABE)


def load_work():
    return [l.split('\t', 1)[1] for l in open(P('work.tsv'), encoding='utf-8').read().split('\n') if l]


def load_out():
    o = []
    if os.path.exists(P('out.tsv')):
        for l in open(P('out.tsv'), encoding='utf-8').read().split('\n'):
            if l == '': continue
            k, s = l.split('\t', 1)
            if int(k) != len(o) + 1: die('out.tsv hỏng thứ tự tại k=' + k)
            o.append(s)
    return o


def save_out(o):
    open(P('out.tsv'), 'w', encoding='utf-8', newline='\n').write(''.join(f'{i + 1}\t{s}\n' for i, s in enumerate(o)))


def load_gl():
    g = []
    if os.path.exists(P('thuat-ngu.tsv')):
        for l in open(P('thuat-ngu.tsv'), encoding='utf-8').read().split('\n'):
            c = l.rstrip('\r').split('\t')
            if len(c) >= 2 and c[0].strip() and not c[0].startswith('#') and c[1].strip():
                g.append((c[0].strip(), c[1].strip(), c[2].strip() if len(c) > 2 else ''))
    return g


def env():
    meta = jl('meta.json')
    if not meta: die('chưa prep (không thấy dp_work/meta.json)')
    codes = re.compile(CODE0 + ('|' + meta['codes'] if meta.get('codes') else ''))
    cs = set(open(P('charset.txt'), encoding='utf-8').read()) if os.path.exists(P('charset.txt')) else None
    return meta, codes, cs


def fitws(src, d):
    return src[:len(src) - len(src.lstrip())] + d.strip() + src[len(src.rstrip()):]


def probs(src, d, lim, cs, codes):
    if not d.strip(): return ['trống']
    r = []
    a, b = Counter(codes.findall(src)), Counter(codes.findall(d))
    if a != b:
        r.append('mã: ' + ' '.join(filter(None, [('thiếu ' + ' '.join((a - b).elements())) if a - b else '',
                                                  ('thừa ' + ' '.join((b - a).elements())) if b - a else ''])))
    if CJK.search(d): r.append('còn chữ Hán/ký tự rộng')
    if lim and len(d) > lim: r.append(f'dài {len(d)}>{lim}')
    if cs is not None:
        miss = {c for c in d.replace(NLE, '').replace(TABE, '') if c not in cs and not c.isspace()}
        if miss: r.append('thiếu glyph: ' + ''.join(sorted(miss)))
    return r


def cmd_prep(a):
    if os.path.exists(P('meta.json')) and not a.force:
        die('đã có dp_work/ - thêm --force để làm lại (xóa out.tsv)')
    os.makedirs(W, exist_ok=True)
    meta, rows, _ = loader(a.file)(a.file, a)
    meta.update(file=os.path.abspath(a.file), lang=a.lang, codes=a.codes or '', enc=a.enc,
                args=dict(src=a.src, dst=a.dst, ctx=a.ctx, lim=a.lim, sheet=a.sheet, enc=a.enc))
    codes = re.compile(CODE0 + ('|' + a.codes if a.codes else ''))
    uniq, pos2k, ctxd, limd, skip, skc, work, cc, chars = {}, {}, {}, {}, [], Counter(), [], Counter(), 0
    for pos, src, ctx, lim in rows:
        cc.update(codes.findall(src))
        if skip_it(src, a.lang):
            skip.append(str(pos)); skc[src] += 1; continue
        if src not in uniq:
            uniq[src] = len(uniq) + 1
            k = uniq[src]
            work.append(f'{k}\t{esc(src)}'); chars += len(src)
            if ctx.strip(): ctxd[str(k)] = ' '.join(ctx.split())
            if lim: limd[str(k)] = lim
        elif lim:
            limd[str(uniq[src])] = min(lim, limd.get(str(uniq[src]), lim))
        pos2k[str(pos)] = uniq[src]
    open(P('work.tsv'), 'w', encoding='utf-8', newline='\n').write('\n'.join(work) + '\n')
    open(P('skipped.tsv'), 'w', encoding='utf-8', newline='\n').write(
        ''.join(f'{n}\t{s}\n' for s, n in skc.most_common()))
    open(P('out.tsv'), 'w').close()
    jd('meta.json', meta); jd('map.json', dict(pos2k=pos2k, ctx=ctxd, lim=limd, skip=skip))
    if a.glossary: shutil.copy(a.glossary, P('thuat-ngu.tsv'))
    elif not os.path.exists(P('thuat-ngu.tsv')):
        open(P('thuat-ngu.tsv'), 'w', encoding='utf-8').write('# gốc\tdịch\tloại(~ = không ép kiểm tra)\tghi chú\n')
    if a.charset:
        open(P('charset.txt'), 'w', encoding='utf-8').write(open(a.charset, encoding='utf-8').read())
    c = meta['cols']
    print(f"định dạng={meta['fmt']} header={meta['hdr']} cột: " + ' '.join(f'{k}={v}' for k, v in c.items() if k != 'dst_new')
          + (' (dst = cột mới)' if c.get('dst_new') else ''))
    print(f'dòng dữ liệu={len(rows)} | cần dịch={len(work)} (duy nhất) | không dịch={len(skip)} ({len(skc)} kiểu) | ký tự gốc={chars}')
    print('mã điều khiển: ' + (' '.join(f'{k}×{v}' for k, v in cc.most_common(15)) or '(không có)'))


def cmd_peek(a):
    w = load_work(); n = min(a.n, len(w))
    for i in sorted({int(j * len(w) / n) for j in range(n)}):
        print(f'{i + 1}\t{w[i][:70]}')


def cmd_next(a):
    env(); w, o, m = load_work(), load_out(), jl('map.json')
    done, N = len(o), len(w)
    if done >= N:
        print(f'XONG {N}/{N} - chạy: dp.py check'); return
    sel, chars, i = [], 0, done
    while i < N and len(sel) < a.n and not (sel and chars + len(w[i]) > a.chars):
        sel.append(i); chars += len(w[i]); i += 1
    print(f'# {done}/{N} -> lô k={sel[0] + 1}..{sel[-1] + 1}')
    if a.vai and os.path.exists(P('vai.md')): print(open(P('vai.md'), encoding='utf-8').read().rstrip())
    blob = '\n'.join(w[j] for j in sel)
    hit = [f'{g}={d}' for g, d, _ in load_gl() if g in blob][:200]
    if hit: print('# thuật ngữ: ' + ' ; '.join(hit))
    last = None
    for j in sel:
        k = str(j + 1)
        if m['ctx']:
            c = m['ctx'].get(k, '-')[:30]
            if c != last: print('#ctx ' + c); last = c
        l = m['lim'].get(k)
        print(f'{k}\t{w[j]}' + (f'\t≤{l}' if l else ''))


def parse_in():
    items = []
    for ln in sys.stdin.read().split('\n'):
        ln = ln.rstrip('\r')
        if not ln.strip(): continue
        m = re.match(r'(\d+)\t(.*)$', ln) or re.match(r'(\d+) (.*)$', ln)
        if not m: die(f'dòng không đúng dạng "k<TAB>dịch": {ln[:50]}')
        if '\t' in m.group(2): die(f'dòng k={m.group(1)} chứa TAB thừa (đừng chép cột ≤giới hạn)')
        items.append((int(m.group(1)), m.group(2)))
    if not items: die('stdin rỗng')
    return items


def cmd_add(a, fix=False):
    meta, codes, cs = env(); w, o, m = load_work(), load_out(), jl('map.json')
    items = parse_in(); done = len(o)
    if fix:
        for k, _ in items:
            if not 1 <= k <= done: die(f'fix: k={k} chưa dịch hoặc không tồn tại')
    else:
        exp = list(range(done + 1, done + 1 + len(items)))
        if [k for k, _ in items] != exp:
            die(f'cần k={exp[0]}..{exp[-1]} liền mạch, nhận {items[0][0]}..{items[-1][0]} (nhảy/trùng/thiếu). Chưa ghi gì.')
        if exp[-1] > len(w): die('vượt tổng số dòng')
    bad = []
    for k, d in items:
        d = fitws(w[k - 1], d)
        if fix: o[k - 1] = d
        else: o.append(d)
        pr = probs(w[k - 1], d, m['lim'].get(str(k)), cs, codes)
        if pr: bad.append(f'⚠ {k}: ' + '; '.join(pr))
    save_out(o)
    for b in bad[:30]: print(b)
    if len(bad) > 30: print(f'... và {len(bad) - 30} cảnh báo nữa')
    print(f'✓ {"sửa" if fix else "+"}{len(items)} | {len(o)}/{len(w)}')


def cmd_check(a):
    meta, codes, cs = env(); w, o, m = load_work(), load_out(), jl('map.json')
    bad = []
    if len(o) < len(w): bad.append(f'CHƯA XONG {len(o)}/{len(w)}')
    pl = [f'{i + 1}: ' + '; '.join(p) for i in range(len(o))
          if (p := probs(w[i], o[i], m['lim'].get(str(i + 1)), cs, codes))]
    gl = [f'{i + 1}: {g}→{d}' for g, d, t in load_gl() if t != '~'
          for i in range(len(o)) if g in w[i] and d.lower() not in o[i].lower()]
    for name, lst in (('dòng lỗi', pl), ('thuật ngữ lệch', gl)):
        if lst:
            bad.append(f'{name}: {len(lst)}'); bad += lst[:40]
            if len(lst) > 40: bad.append(f'... +{len(lst) - 40}')
    print('\n'.join(bad) if bad else f'OK {len(o)}/{len(w)}')


def cmd_sample(a):
    env(); w, o = load_work(), load_out()
    for i in sorted(random.Random(1).sample(range(len(o)), min(a.n, len(o)))):
        print(f'{i + 1}\t{w[i]}\n \t⇒ {o[i]}')


def cmd_merge(a):
    meta, codes, cs = env(); w, o, m = load_work(), load_out(), jl('map.json')
    if len(o) < len(w) and not a.partial:
        die(f'mới dịch {len(o)}/{len(w)} - dịch tiếp, hoặc thêm --partial (dòng chưa dịch giữ nguyên gốc)')
    ns = argparse.Namespace(**meta['args'])
    m2, rows, data = loader(meta['file'])(meta['file'], ns)
    m2.update(enc=meta['enc'])
    fmt = meta['fmt']
    tm = {}
    for pos, src, _, _ in rows:
        k = m['pos2k'].get(str(pos))
        t = o[k - 1] if k and k <= len(o) else None
        if t is None: tm[pos] = src
        else: tm[pos] = t if fmt in ('tsv', 'prop') else t.replace(NLE, '\n').replace(TABE, '\t')
    outp = os.path.join(a.out or P('final'), os.path.basename(meta['file']))
    os.makedirs(os.path.dirname(outp), exist_ok=True)
    c = m2['cols']
    if fmt == 'tsv':
        d, out = c['dst'], list(data)
        if m2['hdr'] and c.get('dst_new'):
            x = data[0].split('\t'); x += [''] * (d + 1 - len(x)); x[d] = 'dịch'; out[0] = '\t'.join(x)
        for pos, t in tm.items():
            x = data[pos].split('\t'); x += [''] * (d + 1 - len(x)); x[d] = t; out[pos] = '\t'.join(x)
        wr(outp, m2['nl'].join(out) + (m2['nl'] if m2['tail'] else ''), m2['enc'], m2['bom'])
    elif fmt == 'prop':
        out = list(data)
        for pos, t in tm.items(): out[pos] = PR.match(data[pos]).group(1) + t
        wr(outp, m2['nl'].join(out) + (m2['nl'] if m2['tail'] else ''), m2['enc'], m2['bom'])
    elif fmt == 'json':
        for pos, t in tm.items(): data[pos] = t
        wr(outp, json.dumps(data, ensure_ascii=False, indent=m2['indent']), m2['enc'], m2['bom'])
    else:
        wb, ws = data
        if m2['hdr'] and c.get('dst_new'): ws.cell(1, c['dst'] + 1).value = 'dịch'
        for pos, t in tm.items(): ws.cell(pos, c['dst'] + 1).value = t
        wb.save(outp)
    print('→ ' + outp)


def main():
    ap = argparse.ArgumentParser(prog='dp.py'); sp = ap.add_subparsers(dest='cmd', required=True)
    p = sp.add_parser('prep'); p.add_argument('file')
    p.add_argument('--lang', default='zh', choices=['zh', 'en'])
    for f in ('codes', 'glossary', 'charset', 'sheet'): p.add_argument('--' + f)
    p.add_argument('--enc', default='utf-8'); p.add_argument('--force', action='store_true')
    for f in HDR: p.add_argument('--' + f, type=int)
    p = sp.add_parser('next'); p.add_argument('-n', type=int, default=300)
    p.add_argument('--chars', type=int, default=5000); p.add_argument('--vai', action='store_true')
    sp.add_parser('add'); sp.add_parser('fix'); sp.add_parser('check')
    for c, n in (('peek', 60), ('sample', 30)):
        p = sp.add_parser(c); p.add_argument('n', nargs='?', type=int, default=n)
    p = sp.add_parser('merge'); p.add_argument('--partial', action='store_true'); p.add_argument('--out')
    a = ap.parse_args()
    {'prep': cmd_prep, 'peek': cmd_peek, 'next': cmd_next, 'add': cmd_add, 'fix': lambda a: cmd_add(a, True),
     'check': cmd_check, 'sample': cmd_sample, 'merge': cmd_merge}[a.cmd](a)


if __name__ == '__main__':
    main()
